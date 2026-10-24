# Kitchen v3 costing: material quantities from the v3 cut list + hardware CSVs, unit prices = Philippine
# market estimates (Oct 2026, low = budget brand / online, high = branded / hardware store). Edit the price
# columns in Kitchen_v3_Costing.xlsx with your supplier quotes; totals are formulas.
# Run:  uv run --no-project --with openpyxl python make_kitchen_costing.py

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import kitchen_cutlist_core as c

OUT = os.path.join(c.OUT_DIR, "Kitchen_v3_Costing.xlsx")
cut_csv, hw_csv = c.csv_paths("v3")
rows = c.read_rows(cut_csv)
layouts, _ = c.layout(rows)
n18, n6 = len(layouts[c.PVC18]), len(layouts.get(c.PVC6, []))
pieces18 = sum(r[3] for r in rows if r[7] == c.PVC18)
ctop_m2 = sum(r[3] * r[4] * r[5] for r in rows if r[7] == c.CTOP) / 1e6
hinges = sum(r[3] * c.hinges_for(max(r[4], r[5])) for r in rows if c.front_kind(r[2]) == "Door")
fronts = sum(r[3] for r in rows if c.front_kind(r[2]) in ("Door", "Drawer"))
edge_m = sum(r[3] * 2 * (r[4] + r[5]) for r in rows if c.front_kind(r[2])) / 1000 * 1.1
hw = c.read_hardware(hw_csv)


def hwq(*keys):
    return sum(c.hw_qty_unit(size, q)[0] for (item, size), (q, _) in hw.items() if all(k in item for k in keys))


runners500 = hwq("Full-extension", "500")
runners450 = hwq("Full-extension", "450")
baskets = sum(q for (item, size), (q, _) in hw.items()
              if any(k in item for k in ("basket", "tray (", "Cutlery")) and "Under-sink" not in item)
led_m = sum(c.hw_qty_unit(size, q)[0] for (item, size), (q, _) in hw.items() if "LED strip" in item)
cab24 = hwq("24 V cable")
cab230 = hwq("230 V feed")

# (category, item, qty, unit, low, high, note)
L = [
    ("Board", "PVC board 18 mm 4x8 (cabinet grade, e.g. Versaply)", n18 + 1, "sheet", 1700, 3600,
     f"{n18} nested + 1 spare; {pieces18} pieces"),
    ("Board", "PVC board 6 mm 4x8", n6, "sheet", 700, 1200, "toe-kick drawer bottoms"),
    ("Countertop", "Granite countertop, fabricated + installed (4 pieces)", round(ctop_m2, 2), "m2", 6500, 15000,
     "high = premium granite / quartz; fabricator measures on site"),
    ("Countertop", "Countertop cut-outs (sink, cooktop, pipe notch)", 3, "set", 850, 1500, ""),
    ("Fittings", "Concealed hinge 35 mm soft-close, clip-on", hinges, "pc", 65, 350, "low = generic, high = Blum / Hafele"),
    ("Fittings", "Handles or push-to-open latches", fronts, "pc", 80, 300, ""),
    ("Runners", "Full-extension soft-close runners 500 mm", runners500, "pair", 380, 930, "Giano ~380, Hafele ~790-930"),
    ("Runners", "Full-extension soft-close runners 450 mm", runners450, "pair", 360, 900, "under-sink tray"),
    ("Runners", "Toe-kick drawer runners 400 mm, push-to-open", hwq("Toe-kick drawer runner"), "pair", 450, 1000, ""),
    ("Racks", "Stainless pull-out baskets / trays (drawers)", baskets, "pc", 1500, 4500, "304 SS; 900-wide ones cost more"),
    ("Racks", "Under-sink pull-out tray", 1, "pc", 1500, 4000, ""),
    ("Racks", "Dish drainer set 3-tier (900 cabinet)", 1, "set", 2500, 7000, "plate + cup tiers + drip tray"),
    ("Racks", "Detergent caddy (soap cabinet)", 1, "pc", 500, 1500, ""),
    ("Fasteners", "Confirmat / coarse-thread screws", round(4 * pieces18, -1), "pc", 1.5, 3, ""),
    ("Fasteners", "Short screws 4x16", 8 * hinges + 12 * (runners500 + runners450 + hwq("Toe-kick drawer runner")),
     "pc", 0.5, 1, ""),
    ("Fasteners", "Shelf pins 5 mm", hwq("Shelf pins"), "pc", 3, 8, ""),
    ("Wall mount", "Wall-cabinet hanging brackets", 2 * c.WALL_CABINETS_REV["v3"], "pc", 150, 400, ""),
    ("Wall mount", "Masonry anchors 8 mm", 4 * c.WALL_CABINETS_REV["v3"], "pc", 10, 25, ""),
    ("Finishing", "PVC edge band 22 mm (optional, fronts)", round(edge_m), "m", 15, 30, ""),
    ("Finishing", "Neutral-cure silicone", 3, "tube", 250, 450, ""),
    ("Finishing", "Paint / wrap for the exposed corner pipe", 1, "lot", 500, 1000, ""),
    ("Hood", "Built-in telescopic range hood 60 cm, ducted", 1, "pc", 8500, 18000, "e.g. Elba ITH 602 ~8,490"),
    ("Hood", "Flat duct 204x60 + 150 round adapter", 1, "set", 1500, 3000, ""),
    ("Gas", "LPG regulator + 1.5 m hose + clamps", 1, "set", 800, 1500, ""),
    ("Gas", "Vent grilles (tank door) + 50 mm hose grommet", 1, "set", 350, 900, "2 grilles + grommet"),
    ("Gas", "Heat-shield sheet, aluminium-faced", 1, "pc", 500, 1200, ""),
    ("Appliance fit", "Vent grille behind induction front", 1, "pc", 200, 500, ""),
    ("Lighting", "24 V COB LED strip + aluminium profile + diffuser", round(led_m, 1), "m", 400, 900, "4 channels"),
    ("Lighting", "24 V LED drivers (20 + 30 + 30 + 60 W)", 4, "pc", 650, 1600, ""),
    ("Lighting", "Wall switches (under-cabinet lights)", 2, "pc", 250, 600, ""),
    ("Lighting", "PIR motion sensors (toe-kick)", 2, "pc", 400, 900, ""),
    ("Lighting", "24 V cable 2 x 0.75 mm2", round(cab24, 1), "m", 25, 50, ""),
    ("Lighting", "230 V cable in conduit", round(cab230, 1), "m", 40, 80, ""),
    ("Lighting", "Cable grommets 10 mm", 8, "pc", 15, 40, ""),
    ("Electrical", "New outlets (2 doubles in the appliance garage, 1 backsplash)", 3, "pc", 300, 700,
     "devices + boxes only"),
    ("Labour", "Carpentry: cutting, assembly, installation (2 carpenters ~15-20 days)", 1, "lot", 30000, 60000,
     "rule of thumb; ask your carpenter for a quote"),
    ("Labour", "Electrician: new circuits, outlets, lighting, island conduit", 1, "lot", 8000, 20000, ""),
]

wb = Workbook()
ws = wb.active
ws.title = "Costing"
B = Font(bold=True)
ws["A1"] = "Kitchen v3 - costing estimate (PHP)"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = ("Quantities from the v3 model. Unit prices are Philippine market estimates (Oct 2026) - yellow cells: "
            "replace with your supplier quotes. Appliances (fridge, stove, washer, induction cooktop, microwave) "
            "are not included.")
hdr = ["Category", "Item", "Qty", "Unit", "Low unit", "High unit", "Low total", "High total", "Note"]
for i, h in enumerate(hdr, 1):
    ws.cell(4, i, h).font = B
yellow = PatternFill("solid", fgColor="FFFF99")
r = 5
for cat, item, q, unit, lo, hi, note in L:
    for col, v in enumerate((cat, item, q, unit, lo, hi), 1):
        ws.cell(r, col, v)
    ws.cell(r, 5).fill = ws.cell(r, 6).fill = yellow
    ws.cell(r, 7, f"=C{r}*E{r}")
    ws.cell(r, 8, f"=C{r}*F{r}")
    ws.cell(r, 9, note)
    for col in (5, 6, 7, 8):
        ws.cell(r, col).number_format = "#,##0"
    r += 1
last = r - 1
ws.cell(r + 1, 2, "Materials + hardware").font = B
ws.cell(r + 2, 2, "Labour").font = B
ws.cell(r + 3, 2, "Contingency 10%").font = B
ws.cell(r + 4, 2, "TOTAL").font = Font(bold=True, size=12)
lab0 = next(i for i, x in enumerate(L) if x[0] == "Labour") + 5
for col in "GH":
    ws[f"{col}{r + 1}"] = f"=SUM({col}5:{col}{lab0 - 1})"
    ws[f"{col}{r + 2}"] = f"=SUM({col}{lab0}:{col}{last})"
    ws[f"{col}{r + 3}"] = f"=ROUND(0.1*({col}{r + 1}+{col}{r + 2}),-2)"
    ws[f"{col}{r + 4}"] = f"=SUM({col}{r + 1}:{col}{r + 3})"
    for k in range(1, 5):
        ws[f"{col}{r + k}"].number_format = "#,##0"
        ws[f"{col}{r + k}"].font = B
for col, w in zip("ABCDEFGHI", (13, 62, 8, 7, 10, 10, 12, 12, 48)):
    ws.column_dimensions[col].width = w
wb.save(OUT)

lo = sum(q * a for _, _, q, _, a, _, _ in L)
hi = sum(q * b for _, _, q, _, _, b, _ in L)
by = {}
for cat, _, q, _, a, b, _ in L:
    x = by.setdefault(cat, [0, 0]); x[0] += q * a; x[1] += q * b
for cat, (a, b) in by.items():
    print(f"{cat:14s} {a:>10,.0f} {b:>10,.0f}")
lab = by["Labour"]
print(f"{'subtotal':14s} {lo:>10,.0f} {hi:>10,.0f}")
print(f"{'+10% cont.':14s} {round(lo * 0.1, -2):>10,.0f} {round(hi * 0.1, -2):>10,.0f}")
print(f"{'TOTAL':14s} {lo + round(lo * 0.1, -2):>10,.0f} {hi + round(hi * 0.1, -2):>10,.0f}")
print("wrote", OUT, "| hinges", hinges, "fronts", fronts, "baskets", baskets, "led m", led_m, "ctop m2", round(ctop_m2, 2))
