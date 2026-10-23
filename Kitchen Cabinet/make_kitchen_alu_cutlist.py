# Cut list + materials workbook, cutting PDF and cutlistoptimizer.com CSV for the ALUMINIUM kitchen.
# Reads the CSVs written by build_kitchen_alu.py (run that in FreeCAD first), then:
#   - bar cutting plan for 2x1 / profile / profile with handle (first-fit decreasing on stock bars)
#   - nests the cladding parts on 4x8 sheets
#   - writes Kitchen_Alu_CutList_Materials.xlsx, kitchen_alu_cutting_diagram.pdf,
#     cutlistoptimizer_alu_cladding.csv
# Run:  python3 make_kitchen_alu_cutlist.py   (needs openpyxl + matplotlib)

import csv
import math
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Rectangle

from kitchen_alu_core import (BAR_KERF, BAR_L, BAR_TRIM, BARS, HW_CSV, KERF, M_CLAD, M_CTOP, M_HANDLE, M_PROF,
                              M_TUBE, OUT_DIR, SPARE_BARS, SHEET_L, SHEET_W, SPEC, TRIM, bar_used, counts, plan_all,
                              read_hardware, read_rows)
from kitchen_cutlist_core import hw_category, hw_qty_unit

XLSX = os.path.join(OUT_DIR, "Kitchen_Alu_CutList_Materials.xlsx")
PDF = os.path.join(OUT_DIR, "kitchen_alu_cutting_diagram.pdf")

rows = read_rows()
hw = read_hardware(HW_CSV)
bars, sheets, where = plan_all(rows)
wall, base, tube_pcs, clad_perim = counts(rows)
info = {r[0]: r for r in rows}
TAG = {M_TUBE: "T", M_PROF: "P", M_HANDLE: "H"}

F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3A5F")
H_FONT = Font(name=F, bold=True, color="FFFFFF")
B_FONT = Font(name=F)
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
NOTE = Font(name=F, italic=True, color="666666", size=9)
INPUT = Font(name=F, color="0000FF")
YELLOW = PatternFill("solid", fgColor="FFFF99")
thin = Side(style="thin", color="BBBBBB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, names, widths):
    for c, (n, w) in enumerate(zip(names, widths), 1):
        cell = ws.cell(row=row, column=c, value=n)
        cell.font, cell.fill, cell.border = H_FONT, H_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[chr(64 + c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, row, col, val, font=B_FONT, fmt=None):
    c = ws.cell(row=row, column=col, value=val)
    c.font, c.border = font, BOX
    if fmt:
        c.number_format = fmt
    return c


def title(ws, t, note):
    ws["A1"], ws["A2"] = t, note
    ws["A1"].font, ws["A2"].font = TITLE, NOTE


wb = Workbook()

# ---- Cut List
cl = wb.active
cl.title = "Cut List"
title(cl, "Aluminium kitchen - Cut List (mm)",
      "Bars at cut length (mitred profiles: long point, see Note). Bar / sheet: T = 2x1, P = profile, "
      "H = profile with handle, C = cladding sheet. 'see model' = notch / cut-back shown in FreeCAD.")
header(cl, 4, ["Ref #", "Module", "Part", "Material", "Qty", "Length", "Width", "Note", "Bar / sheet",
               "Total m", "Area m²"], [7, 26, 28, 18, 6, 9, 9, 34, 16, 9, 9])
first = 5
for i, r in enumerate(rows):
    rr = first + i
    ref, mod, part, mat, q, L, W, note = r
    for c, v in enumerate((ref, mod, part, mat, q, L, W if W != "" else None, note,
                           ", ".join(sorted(set(where.get(ref, []))))), 1):
        put(cl, rr, c, v, fmt="0.0" if c in (6, 7) else None)
    put(cl, rr, 10, f"=E{rr}*F{rr}/1000" if mat in BARS else None, fmt="0.00")
    put(cl, rr, 11, f"=E{rr}*F{rr}*G{rr}/1000000" if mat in (M_CLAD, M_CTOP) else None, fmt="0.000")
last = first + len(rows) - 1
cl.auto_filter.ref = f"A4:K{last}"
RNG = lambda col: f"'Cut List'!${col}${first}:${col}${last}"

# ---- Bar Plan
bp = wb.create_sheet("Bar Plan")
title(bp, "Bar cutting plan - 2x1, profile, profile with handle",
      f"First-fit decreasing. {BAR_TRIM} mm squared off each bar end, {BAR_KERF} mm blade. Cuts as 'Ref #: length'. "
      "Keep offcuts > 300 mm for corner cleats and repairs.")
header(bp, 4, ["Bar", "Material", "Stock mm", "Pieces", "Cuts (Ref #: length)", "Used mm", "Offcut mm"],
       [7, 18, 10, 8, 110, 10, 10])
rr = 5
for mat in BARS:
    for bi, bar in enumerate(bars[mat], 1):
        for c, v in enumerate((f"{TAG[mat]}{bi}", mat, BAR_L[mat], len(bar),
                               ", ".join(f"{ref}: {L:g}" for ref, L in bar), round(bar_used(bar), 1)), 1):
            put(bp, rr, c, v)
        put(bp, rr, 7, f"=C{rr}-F{rr}-{2 * BAR_TRIM}", fmt="0")
        bp.cell(row=rr, column=5).alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1

# ---- Cladding Layout
sl = wb.create_sheet("Cladding Layout")
title(sl, "Cladding nesting on 4x8 sheets (x, y from the sheet corner, mm) - see kitchen_alu_cutting_diagram.pdf",
      f"Sheet {SHEET_L} x {SHEET_W}, {TRIM} mm edge trim, {KERF} mm blade, guillotine cuts.")
header(sl, 4, ["Sheet", "Ref #", "Module", "Part", "X", "Y", "Along 2440", "Along 1220", "Rotated"],
       [8, 7, 26, 30, 8, 8, 11, 11, 9])
rr = 5
for si, placed in enumerate(sheets, 1):
    for ref, x, y, l, w, rot in sorted(placed, key=lambda p: (p[2], p[1])):
        for c, v in enumerate((f"C{si}", ref, info[ref][1], info[ref][2], round(x), round(y), round(l, 1),
                               round(w, 1), "yes" if rot else ""), 1):
            put(sl, rr, c, v)
        rr += 1

# ---- Summary
sm = wb.create_sheet("Summary", 0)
title(sm, "Aluminium kitchen - cut list & materials summary",
      "Basic materials: 2x1 tube, 4 mm cladding, profile, profile with handle. Blue = from the cutting plan "
      "(re-run make_kitchen_alu_cutlist.py after model changes).")
header(sm, 4, ["Item", "Value", "Unit", "Basis"], [40, 12, 10, 60])
lines, cell = [], {}
for mat in BARS:
    lines += [(f"{mat} pieces", f'=SUMIFS({RNG("E")},{RNG("D")},"{mat}")', "pcs", "Cut List"),
              (f"{mat} total length", f'=SUMIFS({RNG("J")},{RNG("D")},"{mat}")', "m", "Cut List"),
              (f"{mat} stock bars", len(bars[mat]), "bars", f"Bar Plan ({BAR_L[mat] / 1000:.2f} m bars)"),
              (f"{mat} yield", "Y", "%", "cut length / stock length")]
lines += [("Cladding pieces", f'=SUMIFS({RNG("E")},{RNG("D")},"{M_CLAD}")', "pcs", "Cut List"),
          ("Cladding part area", f'=SUMIFS({RNG("K")},{RNG("D")},"{M_CLAD}")', "m²", "Cut List"),
          ("Cladding sheets (4x8)", len(sheets), "sheets", "Cladding Layout / PDF"),
          ("Cladding yield", "YC", "%", "part area / sheet area"),
          ("Countertop pieces", f'=SUMIFS({RNG("E")},{RNG("D")},"{M_CTOP}")', "pcs", "fabricator templates")]
for i, (a, b, c, d) in enumerate(lines):
    r_ = 5 + i
    cell[a] = f"B{r_}"
    if b == "Y":
        b = f"=B{r_ - 2}/(B{r_ - 1}*{BAR_L[a.replace(' yield', '')] / 1000})"
    elif b == "YC":
        b = f"=B{r_ - 2}/(B{r_ - 1}*{SHEET_L * SHEET_W / 1e6})"
    put(sm, r_, 1, a)
    put(sm, r_, 2, b, INPUT if isinstance(b, int) else B_FONT,
        "0.0%" if c == "%" else "0.00" if c in ("m", "m²") else "0")
    put(sm, r_, 3, c)
    put(sm, r_, 4, d)

# ---- Materials
mt = wb.create_sheet("Materials", 1)
title(mt, "List of materials - aluminium",
      "Black qty = formula; blue = count from the model or a rule of thumb (see Basis). "
      "Fill unit prices in the yellow column; the total is at the bottom.")
header(mt, 4, ["Category", "Item", "Specification", "Qty", "Unit", "Basis / note", "Unit price", "Amount"],
       [16, 46, 40, 8, 8, 50, 11, 12])
M = []
for mat in BARS:
    M.append(("Aluminium", mat, f"{SPEC[mat]}, {BAR_L[mat] / 1000:.2f} m bars", f"=Summary!{cell[f'{mat} stock bars']}+{SPARE_BARS[mat]}",
              "bars", f"Bar Plan + {SPARE_BARS[mat]} spare for mis-cuts; offcuts cover cleats"))
M.append(("Aluminium", "Cladding", SPEC[M_CLAD], f"=Summary!{cell['Cladding sheets (4x8)']}+1", "sheets",
          "Cladding Layout + 1 spare"))
for r in rows:
    if r[3] == M_CTOP:
        M.append(("Countertop", f"Countertop - {r[1]}", f"{r[5]:.0f} x {r[6]:.0f} x 38, {r[2]}", r[4], "pc",
                  "fabricator measures on site"))
M.append(("Fasteners", 'Blind (pop) rivets 4 mm (5/32") x 12 - frame joints', "aluminium",
          f"=CEILING(4*Summary!{cell['2x1 pieces']},100)", "pcs", "4 per 2x1 piece (2 per end)"))
M.append(("Fasteners", "Rivets / self-drilling screws - cladding to frame", "every ~200 mm",
          math.ceil(clad_perim / 0.2 / 100) * 100, "pcs", f"cladding edge length {clad_perim:.0f} m / 0.2"))
M.append(("Fasteners", 'Internal corner cleats (1" aluminium angle ~40 mm, or 2x1 offcuts)', "frame joints",
          f"=CEILING(2*Summary!{cell['2x1 pieces']},10)", "pcs", "2 per 2x1 piece"))
M.append(("Fasteners", "Panel adhesive / VHB tape for cladding (keeps panels flat)", "", math.ceil(clad_perim), "m",
          "cladding edge length"))
M.append(("Wall mount", "Masonry anchors 8 mm through the 2x1 back rails", "", 4 * wall + 2 * base + 8, "pcs",
          f"{wall} wall cabinets x 4 + {base} base x 2 + 4 framed panels x 2"))
M.append(("Finishing", "Neutral-cure silicone, clear", "sink, counter joints, backsplash", 3, "tubes", "rule of thumb"))
for (item, size), (q, mods) in hw.items():
    q, spec, unit = hw_qty_unit(size, q)
    cat = "Fittings" if "hinge" in item.lower() or "levelling" in item.lower() else hw_category(item, mods)
    M.append((cat, item, spec, q, unit, ", ".join(sorted(set(mods)))))
for i, (cat, item, spec, q, unit, basis) in enumerate(M):
    r_ = 5 + i
    for c, v in enumerate((cat, item, spec), 1):
        put(mt, r_, c, v)
    is_f = isinstance(q, str) and q.startswith("=")
    put(mt, r_, 4, q, B_FONT if is_f else INPUT, "0.0" if isinstance(q, float) else "0")
    put(mt, r_, 5, unit)
    put(mt, r_, 6, basis)
    put(mt, r_, 7, None, INPUT, "#,##0.00").fill = YELLOW
    put(mt, r_, 8, f"=D{r_}*G{r_}", fmt="#,##0.00")
    for c in (2, 6):
        mt.cell(row=r_, column=c).alignment = Alignment(wrap_text=True, vertical="top")
ml = 4 + len(M)
put(mt, ml + 1, 7, "Total", BOLD)
put(mt, ml + 1, 8, f"=SUM(H5:H{ml})", BOLD, "#,##0.00")
mt.auto_filter.ref = f"A4:H{ml}"

wb.save(XLSX)

# ---------------------------------------------------------------- cutting PDF: bar plans + cladding sheets
COL = {M_TUBE: "#b8bec6", M_PROF: "#d6d9de", M_HANDLE: "#8d939b"}
with PdfPages(PDF) as pdf:
    for mat in BARS:
        bl = bars[mat]
        per_page = 12
        for p0 in range(0, len(bl), per_page):
            chunk = bl[p0:p0 + per_page]
            fig, ax = plt.subplots(figsize=(11.69, 8.27))
            for k, bar in enumerate(chunk):
                yb = (len(chunk) - k) * 1.0
                ax.add_patch(Rectangle((0, yb), BAR_L[mat], 0.6, fill=False, lw=1))
                x = BAR_TRIM
                for ref, L in bar:
                    ax.add_patch(Rectangle((x, yb), L, 0.6, facecolor=COL[mat], edgecolor="#333", lw=0.6))
                    if L > 180:
                        ax.text(x + L / 2, yb + 0.3, f"#{ref}\n{L:g}", ha="center", va="center", fontsize=5.5)
                    x += L + BAR_KERF
                ax.text(-60, yb + 0.3, f"{TAG[mat]}{p0 + k + 1}", ha="right", va="center", fontsize=8)
                ax.text(BAR_L[mat] + 60, yb + 0.3, f"offcut {BAR_L[mat] - bar_used(bar) - 2 * BAR_TRIM:.0f}",
                        ha="left", va="center", fontsize=7)
            ax.set_xlim(-400, BAR_L[mat] + 700)
            ax.set_ylim(0.2, len(chunk) + 1)
            ax.set_yticks([])
            ax.set_xticks(range(0, BAR_L[mat] + 1, 500))
            ax.tick_params(labelsize=7)
            ax.set_title(f"{mat} - {SPEC[mat]} - bars {p0 + 1}-{p0 + len(chunk)} of {len(bl)} "
                         f"({BAR_L[mat]} mm stock, {BAR_KERF} mm blade)", fontsize=10)
            pdf.savefig(fig)
            plt.close(fig)
    for si, placed in enumerate(sheets, 1):
        fig, ax = plt.subplots(figsize=(11.69, 8.27))
        ax.add_patch(Rectangle((0, 0), SHEET_L, SHEET_W, fill=False, lw=1.5))
        used = 0
        for ref, x, y, l, w, rot in placed:
            used += l * w
            ax.add_patch(Rectangle((x, y), l, w, facecolor="#dfe3e8", edgecolor="#333", lw=0.8))
            r = info[ref]
            big = min(l, w) > 120
            txt = f"#{ref}\n{r[2][:24]}\n{l:.0f} x {w:.0f}" if big else f"#{ref} {l:.0f}x{w:.0f}"
            ax.text(x + l / 2, y + w / 2, txt, ha="center", va="center", fontsize=6.5 if big else 5,
                    rotation=0 if l >= w or not big else 90)
        ax.set_xlim(-30, SHEET_L + 30)
        ax.set_ylim(-30, SHEET_W + 30)
        ax.set_aspect("equal")
        ax.set_title(f"Cladding sheet C{si} of {len(sheets)} - {SHEET_L} x {SHEET_W} "
                     f"(yield {used / (SHEET_L * SHEET_W):.0%}, kerf {KERF} mm)", fontsize=11)
        ax.set_xticks(range(0, SHEET_L + 1, 244))
        ax.set_yticks(range(0, SHEET_W + 1, 244))
        ax.tick_params(labelsize=7)
        pdf.savefig(fig)
        plt.close(fig)

# ---------------------------------------------------------------- cutlistoptimizer.com import file (cladding)
with open(os.path.join(OUT_DIR, "cutlistoptimizer_alu_cladding.csv"), "w", newline="") as f:
    wr = csv.writer(f)
    wr.writerow(["Length", "Width", "Qty", "Label"])
    for ref, mod, part, mat, q, L, W, note in rows:
        if mat == M_CLAD:
            wr.writerow([f"{max(L, W):g}", f"{min(L, W):g}", q, f"{ref} {mod} {part}"])

print(f"{len(rows)} rows; bars: " + ", ".join(f"{m} {len(bars[m])}" for m in BARS) +
      f"; cladding sheets {len(sheets)}")
print("wrote", XLSX, PDF, "and cutlistoptimizer_alu_cladding.csv")
